"""Проверка структуры отчёта: каждый день оформлен одинаково.

Скрипт читает готовый документ и подтверждает, что для каждого дня
присутствуют заголовки «Выполнение задания N дня» и «Вывод по N дону»,
а также что в тексте нет устаревших или ошибочных утверждений.

Запуск:  python tools/check_report_111.py <путь к документу>
"""
import re
import sys
from pathlib import Path

from docx import Document

#: Утверждения, которые не должны встречаться в отчёте.
FORBIDDEN = (
    "C++",
    "6 дней",
    "0.1 секунд",
)

#: Утверждения, которые должны присутствовать.
REQUIRED = (
    "18 рабочих дней",
    "0,21 секунды",
    "run_app.py",
    "tbl_presets",
)


def collect(document: Document) -> str:
    """Собирает весь текст документа, включая таблицы.

    Args:
        document: Открытый документ Word.

    Returns:
        Текст документа одной строкой-разделителем.
    """
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def check_structure(document: Document) -> list:
    """Проверяет наличие обязательных заголовков по дням.

    Args:
        document: Открытый документ Word.

    Returns:
        Список найденных несоответствий.
    """
    problems = []
    headings = [paragraph.text.strip() for paragraph in document.paragraphs
                if paragraph.style.name.startswith("Heading")]
    for day in range(1, 12):
        task = any(re.search(r"Выполнение задания\s*%d\s*дня" % day, item)
                   for item in headings)
        conclusion = any(re.search(r"Вывод по\s*%d\s*дону" % day, item)
                         for item in headings)
        print("  день %2d | выполнение задания: %-4s | вывод: %s"
              % (day, "да" if task else "НЕТ", "да" if conclusion else "НЕТ"))
        if not task:
            problems.append(f"день {day}: нет «Выполнение задания {day} дня»")
        if not conclusion:
            problems.append(f"день {day}: нет «Вывод по {day} дону»")
    return problems


def check_facts(text: str) -> list:
    """Проверяет отсутствие устаревших утверждений.

    Args:
        text: Полный текст документа.

    Returns:
        Список найденных несоответствий.
    """
    problems = []
    for needle in FORBIDDEN:
        if needle in text:
            problems.append(f"устаревшее утверждение: «{needle}»")
    for needle in REQUIRED:
        if needle not in text:
            problems.append(f"отсутствует утверждение: «{needle}»")
    return problems


def main() -> int:
    """Проверяет документ и печатает вердикт."""
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(
        r"C:\Users\rokin\Downloads\Отчёт_день_1-11.docx")
    if not path.exists():
        print("Файл не найден:", path)
        return 1
    document = Document(path)
    print("Проверка отчёта:", path.name)
    print()
    problems = check_structure(document)
    print()
    text = collect(document)
    print("Проверка фактов:")
    for needle in FORBIDDEN:
        print("  устаревшее «%s»: %s" % (needle, "НЕТ" if needle not in text else "ЕСТЬ"))
    for needle in REQUIRED:
        print("  требуемое «%s»: %s" % (needle, "есть" if needle in text else "НЕТ"))
    problems += check_facts(text)
    print()
    print("Абзацев: %d, таблиц: %d, рисунков: %d"
          % (len(document.paragraphs), len(document.tables),
             len(document.inline_shapes)))
    print()
    if problems:
        print("Найденные несоответствия:")
        for item in problems:
            print("  -", item)
        return 1
    print("Вердикт: отчёт оформлен корректно, все 11 дней на месте.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())